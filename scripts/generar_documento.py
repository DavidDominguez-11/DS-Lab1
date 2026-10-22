"""Genera docs/Documento_de_Diseno.docx (version final) a partir de ../template.docx (caratula y estilos de la plantilla).

Uso:  python scripts/generar_documento.py [ruta/al/template.docx]
Luego exportar el .docx a PDF desde Word para la entrega.
"""
import copy
import sys

import pandas as pd
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, Twips

RAIZ = Path(__file__).resolve().parents[1]
PLANTILLA = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ.parent / "template.docx"
SALIDA = RAIZ / "docs" / "Documento_de_Diseno.docx"
REPO = "https://github.com/DavidDominguez-11/DS-Lab1 (rama lab10)"

d = docx.Document(PLANTILLA)
body = d.element.body
paras = list(d.paragraphs)

# ---- caratula: se conserva tal cual y solo se cambian laboratorio, titulo y repositorio
for r in paras[9].runs:
    if r.text == "X":
        r.text = "10"
paras[10].runs[0].text = "Visualización Interactiva: Documento de diseño"
rep = paras[27].runs[-1]
nuevo = copy.deepcopy(rep._r)
rep._r.addnext(nuevo)
for t in nuevo.findall(qn("w:t")):
    t.text = " " + REPO
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")

# ---- se borra el contenido de ejemplo (todo despues del salto de pagina de la caratula)
plantilla_cuerpo = copy.deepcopy(paras[40]._p)
sect = body.find(qn("w:sectPr"))
borrar, vivo = [], False
for el in list(body):
    if vivo and el is not sect:
        borrar.append(el)
    if el is paras[35]._p:
        vivo = True
for el in borrar:
    body.remove(el)


def _anexar(el):
    sect.addprevious(el)


def h1(txt):
    _anexar(d.add_paragraph(txt, style="Heading 1")._p)


def h2(txt):
    _anexar(d.add_paragraph(txt, style="Heading 2")._p)


def parrafo(txt, negrita_inicio=None):
    """Parrafo justificado con el formato del cuerpo de la plantilla. negrita_inicio: texto en negrita al principio."""
    p = copy.deepcopy(plantilla_cuerpo)
    runs = p.findall(qn("w:r"))
    base = runs[0]
    for r in runs[1:]:
        p.remove(r)
    if negrita_inicio:
        rb = copy.deepcopy(base)
        rb.find(qn("w:rPr")).insert(1, OxmlElement("w:b"))
        rb.find(qn("w:t")).text = negrita_inicio
        rb.find(qn("w:t")).set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        base.addprevious(rb)
        txt = " " + txt
    t = base.find(qn("w:t"))
    t.text = txt
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    _anexar(p)


def tabla(filas, anchos):
    """Tabla con el estilo GridTable1Light3 de la plantilla; anchos en twips (suman ~9350)."""
    t = d.add_table(rows=len(filas), cols=len(filas[0]))
    tblPr = t._tbl.tblPr
    for hijo in list(tblPr):
        tblPr.remove(hijo)
    est = OxmlElement("w:tblStyle")
    est.set(qn("w:val"), "GridTable1Light3")
    ancho = OxmlElement("w:tblW")
    ancho.set(qn("w:w"), "5000")
    ancho.set(qn("w:type"), "pct")
    look = OxmlElement("w:tblLook")
    look.set(qn("w:val"), "0020")
    look.set(qn("w:firstRow"), "1")
    for hijo in (est, ancho, look):
        tblPr.append(hijo)
    for i, g in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
        g.set(qn("w:w"), str(anchos[i]))
    for i, fila in enumerate(filas):
        for j, txt in enumerate(fila):
            celda = t.cell(i, j)
            celda.width = Twips(anchos[j])
            p = celda.paragraphs[0]
            p.paragraph_format.space_before = Pt(1.8)
            p.paragraph_format.space_after = Pt(1.8)
            r = p.add_run(txt)
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
    _anexar(t._tbl)
    sp = d.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)
    _anexar(sp._p)


def figura(archivo, leyenda):
    """Imagen centrada (captura del tablero) con leyenda debajo."""
    ruta = RAIZ / "docs" / "img" / archivo
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ruta), width=Inches(6.2))
    _anexar(p._p)
    c = d.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(leyenda)
    r.font.name = "Calibri"
    r.font.size = Pt(9)
    r.italic = True
    _anexar(c._p)


# =============================================================== contenido
h1("1. Introducción")
parrafo("El Observatorio de Turismo Sostenible de Guatemala es un tablero interactivo dirigido a directivos y tomadores de "
        "decisiones del sector turístico que no necesariamente tienen formación en estadística. Utiliza la base de movimientos "
        "migratorios de viajeros internacionales (161,036 registros, enero 2009 a junio 2026) y reutiliza los modelos de "
        "predicción de los Laboratorios 1 y 2, sin construir modelos nuevos. Este documento describe las decisiones de diseño: "
        "preguntas de gestión, alcance, herramienta, paleta, interfaz, visualizaciones, modelos y planificación.")

h1("2. Preguntas de gestión y preguntas exploratorias")
parrafo("Se seleccionaron las preguntas que pueden responderse con las fuentes disponibles. Cada una determina una página "
        "del tablero y, en conjunto, la navegación: de lo general (qué pasó) a lo específico (cuándo, de dónde, por dónde) y "
        "finalmente a lo prospectivo (qué podría pasar y qué tan confiable es).")
tabla([
    ["#", "Pregunta de gestión", "Dónde se responde"],
    ["P1", "¿Cómo ha evolucionado el ingreso de viajeros y cómo se recuperó tras la pandemia?", "Página 1: barras por año"],
    ["P2", "¿Qué ocurrió en un año concreto: en qué meses y por qué vías ingresaron?", "Página 1: detalle enlazado"],
    ["P3", "¿En qué meses hay temporada alta y baja?", "Página 2: mapa de calor"],
    ["P4", "¿De qué regiones y países provienen los viajeros y por qué fronteras ingresan?", "Página 2: treemap y barras"],
    ["P5", "¿Qué se espera a futuro y qué tan confiable es esa estimación?", "Página 3: real frente a estimado"],
    ["P6", "¿Qué modelo se equivoca menos y en qué series?", "Página 3: comparación de modelos"],
], [700, 5900, 2750])

h1("3. Alcance del tablero")
parrafo("exploración de datos observados (2009 a junio 2026), 10 visualizaciones interactivas con dos pares "
        "enlazados, control del nivel de detalle, selección de los gráficos a ver, resultados de 7 modelos ya construidos "
        "(SARIMA, Holt-Winters, Prophet y seasonal naive del Laboratorio 1; LSTM simple, LSTM apilado y LSTM global con catch22 "
        "del Laboratorio 2), una tabla comparativa de errores y un mapa de calor que resume qué modelo conviene en cada serie.", "Incluye:")
parrafo("modelos nuevos (el laboratorio lo prohíbe), explicaciones causales de las variaciones, datos ambientales de "
        "sostenibilidad (no están en las fuentes disponibles) y predicciones más allá del periodo de prueba.", "No incluye:")
parrafo("la categoría «Viajero» cambió de definición entre 2022 y 2023 y «Cruceristas» deja de reportarse desde 2023, por lo "
        "que la medida comparable es Turista + Excursionista; 2026 llega solo hasta junio; los modelos se entrenaron con datos "
        "hasta marzo de 2021, afectados por la pandemia, y tienden a subestimar la recuperación posterior. El tablero "
        "comunica estos límites junto a los gráficos.", "Limitaciones:")

h1("4. Herramienta seleccionada y justificación")
parrafo("Se eligió Streamlit con Plotly (Python), ambas herramientas gratuitas y de código abierto.")
tabla([
    ["Criterio", "Streamlit + Plotly", "Alternativas"],
    ["Lenguaje del equipo", "Python, el mismo de los Laboratorios 1 y 2: se reutilizan pandas y los resultados sin traducir código.", "R Shiny obligaría a reescribir el flujo."],
    ["Gráficos enlazados", "Plotly permite seleccionar con clic y Streamlit actualiza el gráfico relacionado.", "Dash lo logra con callbacks, con más código."],
    ["Velocidad de desarrollo", "Un solo archivo y pocas líneas por gráfico; adecuado para tres integrantes y pocos días.", "Panel y Voilà tienen menos documentación."],
    ["Interactividad", "Zoom, información al pasar el cursor, filtros, selección de gráficos y nivel de detalle sin recargar.", "No aplica"],
    ["Costo y despliegue", "Gratuito; se ejecuta con «streamlit run» o en Streamlit Community Cloud.", "No aplica"],
], [2000, 5000, 2350])

h1("5. Diseño de la interfaz")
h2("5.1 Paleta de colores")
parrafo("Cada color conserva su significado en toda la aplicación. La paleta se basa en Okabe-Ito, segura para personas con "
        "daltonismo, y el texto oscuro sobre fondo blanco ofrece alto contraste.")
tabla([
    ["Color", "Significado fijo", "Justificación"],
    ["Azul oscuro (#1B4F72)", "Dato observado: lo que ocurrió", "El azul transmite confianza e institucionalidad y distingue lo medido de lo estimado."],
    ["Azul medio (#0072B2)", "Selección del usuario y detalle enlazado", "Mismo matiz que lo observado: indica «la misma información, con más detalle»."],
    ["Naranja (#D55E00)", "Estimación de un modelo y avisos", "Contrasta con el azul y alerta que es una estimación, no un hecho."],
    ["Verde, rosa y gris (#009E73, #CC79A7, #7F7F7F)", "Un color fijo por modelo", "Distinguibles entre sí y de la serie real, también en escala de grises."],
    ["Escala azul secuencial", "Intensidad (mapa de calor y treemap)", "Un solo matiz: más oscuro significa más viajeros."],
], [2700, 2900, 3750])
h2("5.2 Tipografía y comunicación")
parrafo("Se usa fuente sans-serif a 16 px como base, legible en pantalla y proyector. Cada gráfico lleva una etiqueta «Dato "
        "observado» o «Estimación de modelo» para separar hallazgos de estimaciones, una pregunta guía bajo el título y un "
        "texto «Cómo leerlo». Las selecciones sin datos muestran un mensaje claro en lugar de un gráfico vacío.")
h2("5.3 Bosquejo y navegación")
parrafo("La navegación es un menú lateral con tres páginas y filtros globales siempre visibles. La estructura de la página 1 es:")
tabla([
    ["Zona", "Contenido"],
    ["Barra lateral", "Menú de 3 páginas (Panorama general, Explorar a detalle, Predicciones y modelos) y filtros: rango de años, tipo de viajero y vía de ingreso."],
    ["Encabezado", "Título de la página y pregunta guía que explica qué se puede decidir con ella."],
    ["Tarjetas", "Tres indicadores: total de viajeros, año con más viajeros y variación del último año."],
    ["Gráfico A (izquierda)", "Barras de viajeros por año. Un clic en una barra selecciona el año."],
    ["Gráfico B (derecha)", "Detalle del año seleccionado, enlazado al gráfico A, con control de nivel de detalle (mes, trimestre o vía)."],
    ["Pie", "Texto «Cómo leerlo» y avisos de límites de los datos."],
], [2400, 6950])

h1("6. Visualizaciones y elección del tipo de gráfica")
tabla([
    ["#", "Visualización", "Tipo de gráfica", "Por qué este tipo"],
    ["1", "Viajeros por año", "Barras verticales", "Compara totales entre categorías discretas (años). Un clic selecciona el año y actualiza la visualización 2."],
    ["2", "Detalle del año (mes, trimestre o vía)", "Línea o barras", "La línea muestra la secuencia mensual; las barras comparan trimestres o vías. Enlazada a la 1 y con tres niveles de detalle."],
    ["3", "Estacionalidad año por mes", "Mapa de calor", "Muestra dos dimensiones a la vez; revela temporadas altas y la caída de 2020."],
    ["4", "Origen: región y país", "Treemap", "Representa jerarquías con área proporcional; permite pasar de regiones a países (nivel de detalle)."],
    ["5", "Principales países", "Barras horizontales", "Ranking con etiquetas largas legibles; el usuario decide cuántos mostrar."],
    ["6", "Principales fronteras", "Barras horizontales", "Identifica los puntos de ingreso clave."],
    ["7", "Real frente a estimado", "Líneas (real continua, modelos punteados)", "Compara series en el tiempo; lo estimado se distingue por color y trazo."],
    ["8", "Error por modelo", "Barras horizontales", "Ordena los modelos por error (MAE o RMSE). Un clic selecciona el modelo y actualiza la visualización 9."],
    ["9", "Error del modelo seleccionado, por mes o por año", "Barras con línea en cero", "Muestra cuándo y hacia qué lado se equivoca el modelo (azul: subestima; naranja: sobrestima). Enlazada a la 8."],
    ["10", "Modelo por serie (error relativo)", "Mapa de calor", "Compara los 7 modelos en las 8 series en una sola vista; 1.0 es el mejor de la serie."],
], [500, 2500, 2100, 4250])
parrafo("Pares enlazados: 1 con 2 y 8 con 9. En ambos, un clic en el primer gráfico cambia el segundo; también se puede elegir "
        "con un selector, para quien no use el ratón. Los usuarios eligen qué gráficos ver en la página 2 y qué modelos comparar "
        "en la página 3. Las tarjetas, los filtros y la tabla comparativa son elementos complementarios y no se cuentan como "
        "visualizaciones.")

h2("6.1 Capturas del tablero")
parrafo("Las siguientes capturas corresponden al tablero en funcionamiento (Streamlit, ejecutado localmente).")
figura("01_panorama.jpg", "Figura 1. Página 1, Panorama general: barras por año (visualización 1) enlazadas al detalle del año seleccionado (visualización 2), con tarjetas, filtros y nivel de detalle.")
figura("02_estacionalidad.jpg", "Figura 2. Página 2, Explorar a detalle: mapa de calor de estacionalidad por año y mes (visualización 3).")
figura("03_origen.jpg", "Figura 3. Página 2: treemap de origen de los viajeros por región (visualización 4); el selector permite bajar al nivel de país.")
figura("04_real_vs_estimado.jpg", "Figura 4. Página 3, Predicciones y modelos: serie real frente a los modelos elegidos (visualización 7), con tabla comparativa de errores.")
figura("05_error_modelo.jpg", "Figura 5. Página 3: error por modelo (visualización 8) y error del modelo seleccionado mes a mes (visualización 9, enlazada a la 8).")
figura("06_modelo_por_serie.jpg", "Figura 6. Página 3: error relativo de cada modelo en cada serie (visualización 10); más oscuro significa peor.")

h1("7. Modelos incluidos y resultados")
parrafo("Los modelos se reutilizan de los Laboratorios 1 y 2 sin entrenar modelos nuevos. El script "
        "scripts/exportar_resultados.py guarda las predicciones de SARIMA, Holt-Winters, Prophet y seasonal naive (Laboratorio 1) "
        "y scripts/exportar_lstm.py ejecuta el notebook del Laboratorio 2 y guarda las de los LSTM. Los archivos quedan en la "
        "carpeta data y el tablero solo los lee. El entrenamiento cubre de enero 2009 a marzo 2021 (147 meses) y la prueba de "
        "abril 2021 a junio 2026 (63 meses). Las series son el total (Turista + Excursionista), tres vías de ingreso y cuatro "
        "tipos de viajero. Los LSTM simple y apilado se entrenaron solo para el total y la vía terrestre; el LSTM global con "
        "catch22 cubre las 8 series.")
_m = pd.concat([pd.read_csv(RAIZ / "data" / "metricas.csv"), pd.read_csv(RAIZ / "data" / "metricas_lstm.csv")])
_mejor = _m.loc[_m.groupby("serie").RMSE.idxmin()].sort_values("serie")
tabla([["Serie", "Modelo con menor error", "RMSE (viajeros por mes)"]] +
      [[r.serie, r.modelo, f"{r.RMSE:,.0f}"] for r in _mejor.itertuples()], [3600, 3600, 2150])
parrafo("Los errores son del periodo de prueba (abril 2021 a junio 2026). Los modelos se entrenaron con datos hasta marzo de "
        "2021, afectados por la pandemia, por lo que casi todos subestiman la recuperación posterior (sesgo negativo). Estos "
        "resultados indican qué modelo se acercó más a lo ocurrido, no qué modelo será mejor en el futuro.", "Lectura:")
parrafo("(1) La selección automática auto_arima del notebook del Laboratorio 1 no se volvió a ejecutar; se usan los dos SARIMA "
        "manuales y se elige el de menor AIC, por lo que el error de SARIMA puede diferir del notebook. (2) El LSTM simple "
        "reproduce exactamente el resultado del notebook del Laboratorio 2. (3) La versión actual de la librería pycatch22 "
        "renombró varias características respecto a la del notebook; se tradujeron los nombres, y el LSTM global con catch22 "
        "puede diferir un poco del valor publicado en ese notebook.", "Notas de reproducibilidad:")

h1("8. Planificación de tareas")
tabla([
    ["Tarea", "Responsable", "Fecha de término", "Estado"],
    ["Preguntas de gestión, alcance y paleta", "Gabriel Bran", "22 de octubre (avance)", "Hecha"],
    ["Exportación de predicciones y métricas del Laboratorio 1", "Gabriel Bran", "22 de octubre (avance)", "Hecha"],
    ["Páginas 1 y 2 del tablero (barras enlazadas, mapa de calor, treemap)", "David Domínguez", "22 de octubre (avance)", "Hecha"],
    ["Documento de diseño (avance)", "Luis Padilla", "22 de octubre (avance)", "Hecha"],
    ["Exportar predicciones LSTM y catch22 del Laboratorio 2", "Luis Padilla", "23 de octubre", "Hecha"],
    ["Página 3: comparación de modelos, error por mes y mensajes de límites", "Gabriel Bran", "24 de octubre", "Hecha"],
    ["Usabilidad, accesibilidad y prueba con un usuario no técnico", "David Domínguez", "24 de octubre", "Pendiente"],
    ["Documento de diseño final (PDF) y enlace del repositorio", "Todos", "25 de octubre (entrega final)", "En curso"],
], [4300, 1700, 2150, 1200])

h1("9. Cómo ejecutar el tablero")
parrafo("Instalar dependencias con «pip install -r requirements_tablero.txt» y ejecutar «streamlit run tablero/app.py».")

SALIDA.parent.mkdir(exist_ok=True)
d.save(SALIDA)
print("generado", SALIDA)
